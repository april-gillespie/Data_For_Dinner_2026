from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("deliverables/Data_for_Dinner_Project_Charter_and_Data_Plan.docx")

INK = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
RULE = "CBD2DA"
CAUTION = "7A5A00"

BULLET_NUM_ID = None
DECIMAL_ABSTRACT_ID = None


def set_font(run, name="Calibri", size=11, color="000000", bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_borders(table, color=RULE, size=6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    assert sum(widths_dxa) == 9360
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_text(cell, text, bold=False, color="000000", size=9.5, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    set_font(run, size=size, color=color, bold=bold)
    p.paragraph_format.keep_together = True


def add_table(doc, headers, rows, widths_dxa, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for i, label in enumerate(headers):
        set_cell_shading(header.cells[i], LIGHT_GRAY)
        set_cell_text(header.cells[i], label, bold=True, color=INK, size=font_size)
    for row_values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for i, value in enumerate(row_values):
            set_cell_text(row.cells[i], str(value), size=font_size)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    return table


def add_numbering_definitions(doc):
    global BULLET_NUM_ID, DECIMAL_ABSTRACT_ID
    root = doc.part.numbering_part.element

    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in root.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in root.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    next_abstract_id = max(abstract_ids, default=0) + 1
    next_num_id = max(num_ids, default=0) + 1

    def make_abstract(abstract_id, num_format, level_text, font_name=None):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_format)
        lvl.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), level_text)
        lvl.append(text)
        jc = OxmlElement("w:lvlJc")
        jc.set(qn("w:val"), "left")
        lvl.append(jc)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        lvl.append(p_pr)
        if font_name:
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), font_name)
            fonts.set(qn("w:hAnsi"), font_name)
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        first_num = root.find(qn("w:num"))
        if first_num is None:
            root.append(abstract)
        else:
            root.insert(root.index(first_num), abstract)

    def make_num(num_id, abstract_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id_node = OxmlElement("w:abstractNumId")
        abstract_id_node.set(qn("w:val"), str(abstract_id))
        num.append(abstract_id_node)
        root.append(num)

    bullet_abstract_id = next_abstract_id
    DECIMAL_ABSTRACT_ID = next_abstract_id + 1
    BULLET_NUM_ID = next_num_id
    make_abstract(bullet_abstract_id, "bullet", "•", "Calibri")
    make_num(BULLET_NUM_ID, bullet_abstract_id)
    make_abstract(DECIMAL_ABSTRACT_ID, "decimal", "%1.")


def new_numbered_list_id(doc):
    root = doc.part.numbering_part.element
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in root.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    num_id = max(num_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_id_node = OxmlElement("w:abstractNumId")
    abstract_id_node.set(qn("w:val"), str(DECIMAL_ABSTRACT_ID))
    num.append(abstract_id_node)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    num.append(level_override)
    root.append(num)
    return num_id


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)


def add_list_item(doc, text, numbered=False, bold_prefix=None, list_id=None):
    p = doc.add_paragraph()
    if numbered and list_id is None:
        raise ValueError("Numbered list items require a shared list_id")
    apply_num(p, list_id if numbered else BULLET_NUM_ID)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    p.paragraph_format.keep_together = True
    if bold_prefix and text.startswith(bold_prefix):
        first = p.add_run(bold_prefix)
        set_font(first, bold=True)
        rest = p.add_run(text[len(bold_prefix):])
        set_font(rest)
    else:
        run = p.add_run(text)
        set_font(run)
    return p


def add_label_paragraph(doc, label, text, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    label_run = p.add_run(f"{label}: ")
    set_font(label_run, bold=True, color=INK)
    value_run = p.add_run(text)
    set_font(value_run)
    return p


def add_callout(doc, label, text, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.line_spacing = 1.10
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), CALLOUT)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), accent)
    borders.append(left)
    p_pr.append(borders)
    r1 = p.add_run(f"{label}: ")
    set_font(r1, bold=True, color=accent)
    r2 = p.add_run(text)
    set_font(r2, color=INK)
    return p


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_source_item(doc, number, title, organization, url, note):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.10
    lead = p.add_run(f"{number}. {organization}. ")
    set_font(lead, size=9.5)
    add_hyperlink(p, title, url)
    tail = p.add_run(f" {note}")
    set_font(tail, size=9.5, color=MUTED)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, after=6, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    run = p.add_run(text)
    set_font(run, italic=italic)
    return p


def add_page_break(doc):
    doc.add_page_break()


def set_update_fields(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)
    set_font(run, size=9, color=MUTED)


def build_document():
    doc = Document()
    doc.core_properties.title = "Data for Dinner - Project Charter and Data Plan"
    doc.core_properties.subject = "Food insecurity and food access analysis in Alabama"
    doc.core_properties.author = "Data for Dinner Team"
    doc.core_properties.keywords = "food access, food insecurity, Alabama, USDA, data plan"
    doc.settings.odd_and_even_pages_header_footer = True

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    add_numbering_definitions(doc)
    set_update_fields(doc)

    def populate_footer(footer):
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        fp.paragraph_format.space_before = Pt(0)
        fp.paragraph_format.space_after = Pt(0)
        prefix = fp.add_run("Approved scope  |  Page ")
        set_font(prefix, size=9, color=MUTED)
        add_field(fp, "PAGE")

    populate_footer(section.footer)
    populate_footer(section.even_page_footer)

    # First-page masthead: standard_business_brief with a named title override.
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(4)
    kr = kicker.add_run("TEAM-APPROVED SCOPE")
    set_font(kr, size=9.5, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(5)
    tr = title.add_run("Data for Dinner")
    set_font(tr, size=27, color=INK, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    sr = subtitle.add_run("Project Charter and Data Plan")
    set_font(sr, size=15, color=MUTED)

    add_label_paragraph(doc, "Status", "Updated through August 29, 2026, including Sandra Kopecky's prevailing team story workbook", after=2)
    add_label_paragraph(doc, "Prepared", "August 22, 2026", after=2)
    add_label_paragraph(doc, "Next review", "Saturday, August 29, 2026 at 1:00 p.m. Central", after=2)
    add_label_paragraph(doc, "Primary focus", "Food insecurity and food access in Alabama", after=12)

    add_callout(
        doc,
        "Recommended direction",
        "Treat food insecurity as the broader condition and food access as the measurable Alabama lens. Food allergies, race and gender demographic analysis, pesticides, and straight-line distance are outside the active project scope.",
    )

    add_heading(doc, "1. Executive summary", 1)
    add_body(
        doc,
        "This charter combines the August 22 planning discussion with the decisions recorded on August 25. The study begins with brief global and national context, compares Alabama with a project-defined Southeast region, and then examines food access within Alabama at the census-tract level.",
    )
    add_label_paragraph(doc, "Approved main question", "Which Alabama census tracts have the greatest combined burden of low income and limited access to food retailers, and how does Alabama compare with the selected Southeast states and the United States?", after=8)
    add_body(doc, "The August 25 meeting confirmed four controlling decisions:")
    for item in (
        "Keep food insecurity and food access as distinct concepts.",
        "Use the ten-state, project-defined Southeast comparison group.",
        "Use SNAP-authorized food retailers for the primary local analysis.",
        "Use road-network distance at 1 mile urban and 10 miles rural, and exclude straight-line distance.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "2. Project framing", 1)
    add_heading(doc, "Problem statement", 2)
    add_body(
        doc,
        "Food insecurity reflects whether households can consistently obtain enough food with available resources. Food access is one contributing dimension and includes geographic proximity, transportation, affordability, and other household constraints. This project will not treat the two concepts as interchangeable. Instead, it will use official food-insecurity estimates for context and a clearly defined retailer-access measure for the local Alabama analysis.",
    )

    add_heading(doc, "Research subquestions", 2)
    subquestions = (
        "What does the broader global and U.S. food-insecurity context show?",
        "How does Alabama compare with the selected Southeast states on household food insecurity and retailer access?",
        "Where are low-income and low-access census tracts concentrated within Alabama?",
        "Which areas show additional access burdens related to vehicle availability, SNAP participation, children, or older adults?",
    )
    subquestion_list_id = new_numbered_list_id(doc)
    for item in subquestions:
        add_list_item(doc, item, numbered=True, list_id=subquestion_list_id)

    add_heading(doc, "Working definitions", 2)
    add_label_paragraph(doc, "Food insecurity", "A household-level condition in which access to adequate food is limited or uncertain because of insufficient money or other resources.")
    add_label_paragraph(doc, "Food access", "The ability to reach and obtain food. The primary local measure used here captures retailer proximity and tract-level economic conditions, not the full concept.")
    add_label_paragraph(doc, "Low-income/low-access tract", "A census tract meeting a specified USDA low-income test and a specified distance-and-population access threshold.")
    add_label_paragraph(doc, "Project-defined Southeast", "Alabama, Arkansas, Florida, Georgia, Kentucky, Louisiana, Mississippi, North Carolina, South Carolina, and Tennessee. This is a project choice, not a standard federal regional definition.")

    add_heading(doc, "3. Scope boundaries", 1)
    add_heading(doc, "In scope", 2)
    for item in (
        "Food insecurity as the broader social condition.",
        "Geographic and transportation-related retailer access.",
        "Global and U.S. context, a ten-state Southeast comparison, and an Alabama tract-level focus.",
        "Descriptive comparisons, maps, reproducible data preparation, and transparent limitations.",
        "A seven-minute presentation, approximately seven slides, a video, an IEEE-style paper, and a GitHub repository.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "Out of scope for the current project", 2)
    for item in (
        "Food allergies and allergen-specific commodity analysis.",
        "Race and gender demographic analysis.",
        "Straight-line distance metrics.",
        "Pesticide use, exposure, or residue analysis.",
        "Claims that retailer proximity causes food insecurity or health outcomes.",
        "Individual-level prediction or identification of affected people.",
        "A full assessment of food quality, affordability, prices, public transit, food banks, online purchasing, or diet quality unless the scope is explicitly expanded.",
    ):
        add_list_item(doc, item)

    add_callout(
        doc,
        "Archive rule",
        "Existing allergy and gluten exploration should be preserved in an out-of-scope archive. It should not feed the active variables, analysis, figures, or narrative.",
        accent=CAUTION,
    )

    add_heading(doc, "4. Approved data plan", 1)
    add_body(doc, "The project uses a small, purposeful data stack. Each layer has one job; no source is included merely because it is available.")

    data_rows = [
        ("Global context", "FAOSTAT Suite of Food Security Indicators", "Country or region; annual", "Use one contextual food-insecurity indicator. Do not join it to U.S. tract data."),
        ("U.S. and Southeast outcome", "USDA Current Population Survey Food Security Supplement", "National and state; three-year averages", "Compare household food-insecurity prevalence across the ten states and United States."),
        ("Core local analysis", "USDA 2025 SNAP-authorized Retailer Access Map", "2020-based census tract", "Primary source for Alabama mapping and consistent geographic comparisons."),
        ("Alabama modeled outcome context", "Feeding America Map the Meal Gap 2026", "County and state; observation year 2024", "Report individual-level modeled food-insecurity and food-cost estimates as a separate layer."),
        ("Team story synthesis", "Sandra Kopecky, WomenInData Data Stats and Summary - Story", "Global to Alabama narrative", "Use as the prevailing presentation-development workbook; credit Sandra and validate exact claims against the reproducible tables."),
        ("Optional county context", "USDA Food Environment Atlas", "County and state; varying years", "Use only when a specific county-level indicator fills a documented gap."),
        ("Map geometry", "USDA map service or compatible Census tract geometry", "Census tract", "Match the 2025 access data's 2020 tract basis; record the geography vintage."),
    ]
    add_table(doc, ["Story layer", "Source", "Geographic grain", "Planned use"], data_rows, [1440, 2340, 1980, 3600], font_size=9)
    source_note = doc.add_paragraph("Sources and access links are listed in Appendix B. Data vintages must be recorded at download time. Request-only raw packages stay outside GitHub unless redistribution terms are explicit.")
    source_note.paragraph_format.space_before = Pt(4)
    source_note.paragraph_format.space_after = Pt(4)
    for run in source_note.runs:
        set_font(run, size=9, color=MUTED, italic=True)

    add_heading(doc, "Primary access measure", 2)
    add_callout(
        doc,
        "Primary measure",
        "Use the network-based low-income/low-access definition at 1 mile for urban tracts and 10 miles for rural tracts: at least 500 people or 33 percent of the population live beyond the threshold from the nearest SNAP-authorized food retailer.",
    )
    add_body(
        doc,
        "This measure is current and road-based, but it includes multiple SNAP-authorized store types, including convenience and dollar stores. The paper and presentation must therefore say 'SNAP-authorized food retailer access,' not 'healthy grocery access.' If the team wants the latter construct, it should deliberately switch to the older large-retailer/supermarket measure.",
    )

    add_heading(doc, "Sensitivity checks", 2)
    for item in (
        "Alternative urban threshold: 0.5 mile, with 10 miles for rural tracts.",
        "Alternative rural threshold: 20 miles, with 1 mile for urban tracts.",
        "Straight-line distance is excluded from the analysis and published results.",
        "Reconcile the final Alabama primary result to 20.2 percent before publication.",
        "Separate descriptive comparison with the 2019 large-retailer measure only if clearly labeled; never present 2019 and 2025 as a direct trend.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "Minimal variable whitelist", 2)
    variable_groups = (
        ("Geography", "State, county, census tract, name fields, FIPS/GEOID, and urban/rural status."),
        ("Denominators", "Total population, housing units, and the relevant subgroup population or household count."),
        ("Primary outcome", "Low-income/low-access indicator plus the count and share of people beyond the selected distance threshold."),
        ("Access burden", "No-vehicle/low-access households and SNAP-receiving housing units with limited access."),
        ("Optional subgroups", "Children and adults age 65 or older, reported only with matching source denominators. Race and gender are outside the current scope."),
        ("Provenance", "Source release, data vintage, retrieval date, original field name, transformation, and denominator."),
    )
    for label, text in variable_groups:
        add_label_paragraph(doc, label, text)

    add_heading(doc, "Data-source acceptance criteria", 2)
    for item in (
        "Directly measures the stated construct or supplies a necessary denominator or geography.",
        "Covers the required geography with stable join keys.",
        "Has a documented vintage, definition, source organization, and reproducible access path.",
        "Has acceptable completeness and does not require unsupported imputation.",
        "Can be processed within the team's schedule and presentation constraints.",
        "Adds analytical value that is not already supplied by a retained source.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "5. Acquisition and analysis workflow", 1)
    workflow = (
        "Inventory candidate sources. Record the official page, vintage, geographic grain, file format, join keys, relevant variables, limitations, and keep/drop rationale.",
        "Run an Alabama pilot. Download only the core source and verify that the proposed fields, tract identifiers, and denominators are usable.",
        "Perform quality checks. Test keys, row counts, missingness, range constraints, denominators, group quarters, and geographic coverage.",
        "Freeze the approved snapshot. Preserve the original file, download date, source link, checksum, and selected-variable manifest.",
        "Create reproducible processed tables. Standardize names and geography keys without overwriting raw files.",
        "Produce descriptive comparisons. Build the Southeast state comparison, Alabama tract map, county rollups, and selected subgroup summaries.",
        "Run sensitivity checks and document limitations. Confirm whether conclusions change under alternate thresholds or exclusions.",
        "Export publication-ready outputs. Generate the paper tables, slide figures, captions, and source notes from the same analysis pipeline.",
    )
    workflow_list_id = new_numbered_list_id(doc)
    for item in workflow:
        add_list_item(doc, item, numbered=True, list_id=workflow_list_id)

    add_heading(doc, "6. Quality and interpretation guardrails", 1)
    guardrails = [
        ("Leading zeros", "Store FIPS and GEOID values as text; Alabama's state code is 01."),
        ("Geographic alignment", "Use compatible tract boundaries and explicitly record whether each source is based on 2010 or 2020 tracts."),
        ("Duplicate keys", "Require one row per expected dataset-geography-year key before joining."),
        ("Denominators", "Verify that subgroup and low-access counts do not exceed their corresponding totals."),
        ("Aggregation", "Recompute county and state shares from summed counts. Do not average tract percentages."),
        ("Missing values", "Distinguish missing, suppressed, not applicable, and true zero values; never silently coerce missing data to zero."),
        ("Group quarters", "Review or separately report tracts dominated by dormitories, prisons, military quarters, or institutional housing."),
        ("Survey estimates", "Treat state food-insecurity rates as estimates with sampling uncertainty, not exact population values."),
        ("Causality", "Use descriptive language such as 'associated with' or 'co-occurs with'; do not claim that retailer distance causes insecurity."),
        ("Retailer interpretation", "State that SNAP retailer access does not measure price, quality, product availability, transit, food banks, or individual experience."),
    ]
    add_table(doc, ["Risk", "Required guardrail"], guardrails, [1800, 7560], font_size=9.25)

    add_heading(doc, "7. Team collaboration and capture", 1)
    add_heading(doc, "System of record", 2)
    add_label_paragraph(doc, "Google Drive", "Canonical location for meeting notes, raw-data snapshots, the data inventory, paper drafts, slides, video files, and shared review artifacts.")
    add_label_paragraph(doc, "GitHub", "Canonical location for finalized code, small processed datasets where appropriate, documentation, source manifests, data dictionaries, and reproducible publication outputs. Partial drafts and exploratory files remain outside the published repository.")
    add_label_paragraph(doc, "Email", "Notification channel only. Decisions made by email must be copied into the shared decision log.")
    add_label_paragraph(doc, "Gemini meeting notes", "Draft input only. Within 24 hours, a team member should convert the notes into decisions, open questions, actions, owners, and due dates.")

    add_heading(doc, "Recommended Google Drive structure", 2)
    folder_text = (
        "00_Admin_and_Meetings\n"
        "01_Question_and_Scope\n"
        "02_Data_Inventory\n"
        "03_Data_Raw_Frozen\n"
        "04_Data_Processed\n"
        "05_Analysis_and_QA\n"
        "06_Figures\n"
        "07_Paper_Slides_Video\n"
        "99_Archive_Out_of_Scope"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    for idx, line in enumerate(folder_text.splitlines()):
        if idx:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, name="Consolas", size=9.5, color=INK)

    add_heading(doc, "Data inventory workbook", 2)
    inventory_rows = [
        ("Datasets", "Source, URL, owner, download date, vintage, grain, coverage, keys, format, status, keep/drop rationale, filename, checksum."),
        ("Variables", "Source field, readable name, definition, unit, denominator, year, transformation, planned use, missing-value rule."),
        ("QA", "Row counts, duplicates, missingness, range checks, denominator checks, geography coverage, reconciliation results."),
        ("Decisions", "Date, question, options considered, decision, rationale, approver, and downstream effect."),
        ("Actions", "Task, owner, due date, status, dependency, and completion evidence."),
    ]
    add_table(doc, ["Tab", "Required contents"], inventory_rows, [1800, 7560], font_size=9.25)

    add_heading(doc, "Suggested GitHub structure", 2)
    for item in (
        "README.md - question, scope, methods, setup, source citations, and result summary.",
        "docs/ - charter, decisions, data dictionary, methods, and limitations.",
        "src/ - reusable acquisition, cleaning, QA, analysis, and figure code.",
        "notebooks/ - clearly named exploration and analysis notebooks.",
        "data/ - small derived data only; raw downloads should be ignored and recreated from scripts or a manifest.",
        "outputs/, paper/, slides/ - versioned publication artifacts where file size permits.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "8. Milestones and decision gates", 1)
    milestone_rows = [
        ("August 25, 8:30 p.m.", "Scope approval complete", "Question, comparison region, road-network threshold, exclusions, repository policy, and role placement recorded."),
        ("August 29, 1:00 p.m.", "Team working session", "Finalize the problem statement, review approved data, confirm sensitivity numbers, and begin final visualizations."),
        ("Before full analysis", "Reproducibility gate", "Frozen raw snapshot, source manifest, data dictionary, scripts, and successful QA checks."),
        ("Weekend ending September 13", "Publication and submission", "Rebuilt figures, report/source reconciliation, seven-minute rehearsal, video check, final repository review, recording, and submission."),
    ]
    add_table(doc, ["Date", "Gate", "Required evidence"], milestone_rows, [2160, 2160, 5040], font_size=9.25)

    add_heading(doc, "9. Deliverable plan", 1)
    add_heading(doc, "Seven-slide narrative", 2)
    slide_outline = (
        "Problem, definitions, and main research question.",
        "Global and U.S. food-insecurity context.",
        "Alabama compared with the project-defined Southeast.",
        "Alabama low-income/low-access census-tract map.",
        "Who and where: vehicle access, SNAP households, children, or older adults.",
        "Interpretation, limitations, and what the measure does not capture.",
        "Conclusion, implications, and future work.",
    )
    slide_list_id = new_numbered_list_id(doc)
    for item in slide_outline:
        add_list_item(doc, item, numbered=True, list_id=slide_list_id)

    add_heading(doc, "Paper and video alignment", 2)
    add_body(doc, "The paper, slides, and video should use the same approved definitions, source vintages, numbers, figures, captions, and limitations. The analysis pipeline should create the tables and figures used by all three deliverables so that revisions do not produce conflicting results.")

    add_heading(doc, "10. Decisions recorded on August 25", 1)
    decisions = (
        "Retain the main question and keep food security distinct from food access.",
        "Retain the ten-state comparison and label it the project-defined Southeast.",
        "Use current SNAP-authorized retailer access for the primary local analysis.",
        "Use the road-network 1-mile urban / 10-mile rural primary measure.",
        "Exclude straight-line distance and vary road-network thresholds for sensitivity.",
        "Exclude food allergies, race, and gender from the current analysis.",
        "Retain the validated Feeding America 2024 county and state modeled estimates as a separate outcome layer.",
        "Publish only finalized repository files and place roles in supporting material at the end.",
        "Record and submit the final project during the weekend ending September 13, 2026.",
    )
    for item in decisions:
        add_list_item(doc, item)

    add_heading(doc, "11. Publication and source gate", 1)
    add_callout(
        doc,
        "What may be published",
        "Only reviewed, reproducible, publication-ready files belong in GitHub. Raw downloads, local previews, partial drafts, and exploratory allergen outputs remain outside the published repository.",
    )
    add_body(doc, "A new dataset may be integrated when all of the following are true:")
    for item in (
        "The exact publisher page and downloadable file are recorded.",
        "Release, observation, and retrieval dates are distinguished.",
        "Geography, unit, denominator, identifiers, and missing-value rules are documented.",
        "Access and redistribution terms are documented; raw package publication requires explicit permission.",
        "Processing is reproducible and the source adds a unique analytical purpose.",
        "QA and reconciliation checks pass for Alabama and the project region.",
    ):
        add_list_item(doc, item)

    add_heading(doc, "Appendix A. Starter data inventory", 1)
    add_body(doc, "Create these records before downloading additional data. Status values should be Proposed, Piloting, Accepted, Rejected, or Archived.")
    starter_rows = [
        ("D01", "FAOSTAT Food Security Indicators", "Global context", "Proposed", "Select one food-insecurity indicator and a common comparison period."),
        ("D02", "USDA CPS Food Security Supplement", "U.S./state outcome", "Proposed", "Use published 2022-2024 state averages and record uncertainty."),
        ("D03", "USDA 2025 SNAP-authorized Retailer Access Map", "Primary local analysis", "Piloting", "Test Alabama network-based variables, keys, counts, and geography."),
        ("D04", "USDA Food Environment Atlas", "Optional county context", "Proposed", "Retain only if a documented county-level gap remains."),
        ("D05", "Compatible census-tract geometry", "Mapping", "Proposed", "Use the USDA service or confirmed 2020-based tract geometry."),
        ("D06", "Feeding America Map the Meal Gap 2026", "Alabama county and state outcome context", "Accepted", "Use selected 2024 fields separately; keep the request-only raw archive outside GitHub and retain the reconciliation note."),
        ("D07", "Sandra Kopecky story workbook", "Prevailing team narrative synthesis", "Accepted", "Preserve the submitted workbook unchanged, credit Sandra, and validate exact claims against repository tables."),
        ("A01", "Allergy/gluten exploration", "Out-of-scope archive", "Archived", "Preserve existing work; exclude it from active data and figures."),
    ]
    add_table(doc, ["ID", "Dataset", "Purpose", "Status", "Next action"], starter_rows, [720, 2460, 1740, 1260, 3180], font_size=8.75)

    add_heading(doc, "Appendix B. Official sources", 1)
    add_source_item(
        doc,
        1,
        "Food Access Research Atlas - Download the Data",
        "U.S. Department of Agriculture, Economic Research Service",
        "https://www.ers.usda.gov/data-products/food-access-research-atlas/download-the-data",
        "Accessed August 22, 2026. Includes the 2025 SNAP-authorized Retailer Access Map and 2019 Large Retailer Access Map downloads.",
    )
    add_source_item(
        doc,
        2,
        "SNAP-authorized Retailer Access Map Reference Guide",
        "U.S. Department of Agriculture, Economic Research Service",
        "https://www.ers.usda.gov/data-products/food-access-research-atlas/documentation/snap-authorized-retailer-access-map-reference-guide",
        "Definitions, thresholds, distance types, population groups, and source vintages for the 2025 tract data.",
    )
    add_source_item(
        doc,
        3,
        "Food Access Research Atlas - Documentation, Strengths, and Limitations",
        "U.S. Department of Agriculture, Economic Research Service",
        "https://www.ers.usda.gov/data-products/food-access-research-atlas/documentation",
        "Scope, methods, retailer coverage, geographic alignment, and interpretation limitations.",
    )
    add_source_item(
        doc,
        4,
        "Household Food Security in the United States in 2024",
        "U.S. Department of Agriculture, Economic Research Service",
        "https://ers.usda.gov/sites/default/files/_laserfiche/publications/113623/ERR-358.pdf",
        "National estimates and 2022-2024 state averages, including Table 5.",
    )
    add_source_item(
        doc,
        5,
        "Suite of Food Security Indicators",
        "Food and Agriculture Organization of the United Nations",
        "https://data.fao.org/catalog/dataset/955d6564-40a9-48b4-b51b-f19d65bb3539",
        "Global and national food-security indicators for contextual analysis.",
    )
    add_source_item(
        doc,
        6,
        "Food Environment Atlas - Data Access and Documentation",
        "U.S. Department of Agriculture, Economic Research Service",
        "https://www.ers.usda.gov/data-products/food-environment-atlas/data-access-and-documentation-downloads",
        "Optional county- and state-level food-environment indicators with varying source years.",
    )
    add_source_item(
        doc,
        7,
        "Map the Meal Gap 2026 Report",
        "Feeding America National Organization",
        "https://www.feedingamerica.org/research/map-the-meal-gap/overall-executive-summary",
        "Published July 28, 2026; modeled local food-insecurity and food-cost estimates for observation year 2024.",
    )
    add_source_item(
        doc,
        8,
        "Map the Meal Gap Methodology",
        "Feeding America National Organization",
        "https://www.feedingamerica.org/research/map-the-meal-gap/how-we-got-the-map-data",
        "Methods, geographic models, food budget shortfall, and meal-cost definitions.",
    )
    add_source_item(
        doc,
        9,
        "Map the Meal Gap Data Request",
        "Feeding America National Organization",
        "https://www.feedingamerica.org/research/map-the-meal-gap/by-county",
        "Official request workflow for the source package. The raw archive remains outside GitHub because it contains no explicit redistribution license.",
    )

    limitations_heading = add_heading(doc, "Appendix C. Known USDA measurement limitations to carry forward", 1)
    limitations = (
        "The retailer-access atlas measures an area-based concept at the census-tract level; it does not directly identify individual household experience.",
        "Geographic proximity is only one dimension of access and does not capture affordability, store quality, product availability, public transit, individual mobility, food pantries, or online purchasing.",
        "The 2025 SNAP-retailer measure includes many retailer types but excludes non-SNAP stores, farmers markets, and delivery routes.",
        "Store locations, road networks, ACS periods, census population patterns, and retailer vintages may not align perfectly in time.",
        "The 2019 large-retailer and 2025 SNAP-retailer products use different store universes, methods, and tract bases and should not be interpreted as a simple time series.",
    )
    limitation_paragraphs = [add_list_item(doc, item) for item in limitations]
    limitations_heading.paragraph_format.keep_with_next = True
    for paragraph in limitation_paragraphs[:-1]:
        paragraph.paragraph_format.keep_with_next = True

    roles_heading = add_heading(doc, "Appendix D. Project roles", 1)
    roles_heading.paragraph_format.page_break_before = True
    add_body(doc, "Roles are placed at the end so readers encounter the project content before team assignments.")
    role_rows = [
        ("Sharon Brooks", "Team lead and project management", "Coordinate meetings and review the repository framework."),
        ("Sandra Kopecky", "Data, database, and story synthesis", "Pull and organize data; author the prevailing team story workbook; develop regional framing and Alabama narrative content."),
        ("April Gillespie", "Repository and analysis", "Manage repository updates, technical implementation, documentation, and integration of validated team data."),
        ("Shared", "Final review and delivery", "Validate final numbers, record the presentation, and submit during the weekend ending September 13, 2026."),
    ]
    add_table(doc, ["Team member", "Role", "Responsibility"], role_rows, [1800, 2400, 5160], font_size=9)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path.resolve())
